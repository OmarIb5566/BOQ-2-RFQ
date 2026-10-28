"""Auto-detection of RFQ field values from the uploaded BOQ — only when it is (close to) certain.

A project field (project name / location / owner) is AUTO-FILLED only when all of these hold:
  1. it is introduced by an explicit label in the BOQ title area, page header or footer
     ("Project:", "Client:", "المالك:", "مشروع …"),
  2. the value is copied verbatim from the file,
  3. every place that states it agrees (no two different values),
  4. the local LLM, reading the same lines independently, extracts the same value
     (skipped — with a note — when the LLM is off).
Anything weaker (a value the LLM sees but no label supports, conflicting values, LLM disagreement)
is offered as a SUGGESTION that the user can pick; it is never filled silently. Every auto-filled
value still needs the user's approval before an RFQ is generated.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from ..models import BOQDocument
from ..text_utils import norm

_HF_CODES = re.compile(r'&"[^"]*"|&K[0-9A-Fa-f]{6}|&\d+|&[A-Za-z&]')
_SEP = ":：-–—"


@dataclass
class Evidence:
    text: str          # verbatim text block (may span lines)
    where: str         # "Div03 page header", "Sheet!F8", "PDF page 1"


@dataclass
class Detection:
    value: str = ""
    evidence: str = ""
    source: str = ""                                     # where it was found, for the user ("Div03 page header")
    auto: bool = False                                   # True = fill it (still needs approval)
    note: str = ""
    suggestions: list[tuple[str, str]] = field(default_factory=list)   # (value, why)


# ------------------------------------------------------------------ evidence collection
def collect_evidence(doc: BOQDocument, max_pdf_pages: int = 3) -> list[Evidence]:
    """Title area above the BOQ table + page headers/footers of EVERY BOQ file, as verbatim text blocks."""
    out: list[Evidence] = []
    for f in doc.file_list():
        tag = f"{f.label} · " if doc.multi_file else ""
        try:
            part = _excel_evidence(doc, f) if f.kind == "excel" else _pdf_evidence(doc, f, max_pdf_pages)
        except Exception:  # noqa: BLE001 - one unreadable file must not stop the others
            continue
        out += [Evidence(e.text, tag + e.where) for e in part]
    return out


def _excel_evidence(doc: BOQDocument, f) -> list[Evidence]:
    import openpyxl

    wb = openpyxl.load_workbook(f.work_path, data_only=True)
    rmap = doc.row_map()
    out: list[Evidence] = []
    for sc in doc.scopes:
        if sc.file != f.id or sc.name not in wb.sheetnames:
            continue
        ws = wb[sc.name]
        hdr = [rmap[u].ref.row for u in sc.header_row_uids if u in rmap]
        last = max(hdr) if hdr else 12
        hidden = set(sc.hidden_rows)
        for r in range(1, last + 1):
            if r in hidden:
                continue
            for c in range(sc.first_col, sc.last_col + 1):
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.strip():
                    out.append(Evidence(v, f"{sc.name.strip()}!{ws.cell(r, c).coordinate}"))
        for kind, hf in (("page header", ws.oddHeader), ("page footer", ws.oddFooter),
                         ("first-page header", ws.firstHeader), ("first-page footer", ws.firstFooter)):
            for part in (hf.left, hf.center, hf.right):
                t = _HF_CODES.sub("", part.text or "") if part is not None else ""
                if t.strip():
                    out.append(Evidence(t, f"{sc.name.strip()} {kind}"))
    return out


def _pdf_evidence(doc: BOQDocument, f, max_pages: int) -> list[Evidence]:
    import pymupdf as fitz

    from ..ingest.pdf_reader import _fix_reversed, _join, _lines, _words_from_chars

    sc = doc.scope("pdf", f.id)
    pdf = fitz.open(f.work_path)
    pages = [_words_from_chars(pdf[i]) for i in range(min(max_pages, len(pdf)))]
    _fix_reversed(pages)
    out = []
    for pno, words in enumerate(pages):
        h = pdf[pno].rect.height
        hb = (sc.page_headers.get(pno) or sc.page_headers.get(str(pno))) if sc else None
        top = hb[1] if hb else 0.18 * h
        for ln in _lines(words):
            y0, y1 = min(w.y0 for w in ln), max(w.y1 for w in ln)
            if y1 <= top + 1 or y0 >= 0.93 * h:
                out.append(Evidence(_join(ln), f"PDF page {pno + 1} {'header' if y1 <= top + 1 else 'footer'}"))
    return out


# ------------------------------------------------------------------ label rules
class LabelRules:
    def __init__(self, cfg: dict):
        self.rules: list[tuple[str, str, str]] = []      # (key, normalised label, kind)
        for key, spec in cfg.get("labels", {}).items():
            for lab in spec.get("with_colon", []):
                self.rules.append((key, norm(lab), "colon"))
            for lab in spec.get("prefix_word", []):
                self.rules.append((key, norm(lab), "prefix"))
        self.rules.sort(key=lambda r: -len(r[1]))         # longest label first ("project location" > "project")

    def match(self, line: str) -> Optional[tuple[str, str]]:
        """(key, verbatim value) if the line starts with a known label."""
        orig = re.sub(r"^[\W_]+", "", line.strip())
        n = norm(orig)
        for key, lab, kind in self.rules:
            if kind == "colon" and re.match(rf"{re.escape(lab)}\s*[{re.escape(_SEP)}]", n):
                i = next((k for k, ch in enumerate(orig) if ch in _SEP), None)
                return (key, orig[i + 1:]) if i is not None else None
        for key, lab, kind in self.rules:
            if kind == "prefix" and re.match(rf"{re.escape(lab)}\s+\S", n):
                return key, orig.split(None, 1)[1]
        return None


def _clean(v: str) -> str:
    v = re.sub(r"\s+", " ", v).strip()
    return v.strip(" ,;:-–—،")


def _continues(v: str) -> bool:
    v = v.rstrip()
    return not v or v[-1] in ",-–&،" or v.lower().endswith(" and")


def rule_candidates(evidence: list[Evidence], rules: LabelRules) -> dict[str, list[tuple[str, str]]]:
    found: dict[str, list[tuple[str, str]]] = {}
    for ev in evidence:
        lines = ev.text.split("\n")
        for i, line in enumerate(lines):
            m = rules.match(line)
            if not m:
                continue
            key, val = m
            j, extra = i + 1, 0
            while _continues(val) and j < len(lines) and extra < 2:    # "Project: Bamboo," / "Family House"
                nxt = lines[j].strip()
                j += 1
                if not nxt:
                    break
                if rules.match(nxt):
                    break
                val, extra = f"{val} {nxt}", extra + 1
            val = _clean(val)
            if 2 <= len(val) <= 150 and re.search(r"[^\W\d_]", val):
                found.setdefault(key, []).append((val, ev.where))
    return found


def _same(a: str, b: str) -> bool:
    na, nb = (re.sub(r"[\W_]+", " ", norm(x)).strip() for x in (a, b))
    if not na or not nb:
        return False
    if na == nb:
        return True
    short, long_ = sorted((na, nb), key=len)
    return short in long_ and len(short) / len(long_) >= 0.7


# ------------------------------------------------------------------ main entry
LLM_TASK = """Below are the title, page-header and page-footer lines of a construction BOQ (one per line, \
with their location). Copy EXACTLY (verbatim, same language) the project name, the project location and the \
project owner / client / employer — but ONLY if the text explicitly states it. Use "" for anything that is not \
explicitly stated. Do not guess, translate or complete.

{lines}"""

LLM_SCHEMA = {
    "type": "object",
    "properties": {"project_name": {"type": "string"}, "project_location": {"type": "string"},
                   "owner": {"type": "string"}},
    "required": ["project_name", "project_location", "owner"],
}


def detect_project_fields(doc: BOQDocument, cfg: dict, llm=None, system: str = "",
                          log: Callable[[str], None] = print) -> dict[str, Detection]:
    """Detections keyed by detect-key (project_name / project_location / owner)."""
    evidence = collect_evidence(doc)
    cands = rule_candidates(evidence, LabelRules(cfg))
    llm_ans: Optional[dict] = None
    if llm is not None and evidence:
        lines, seen = [], set()
        for ev in evidence:
            t = re.sub(r"\s+", " ", ev.text).strip()
            if t and t not in seen:
                seen.add(t)
                lines.append(f"[{ev.where}] {t[:200]}")
        try:
            llm_ans = llm.chat_json(system, LLM_TASK.format(lines="\n".join(lines[:80])), LLM_SCHEMA)
        except Exception as e:  # noqa: BLE001
            log(f"RFQ field cross-check: LLM unavailable ({e}).")
    all_text = norm(" ".join(ev.text for ev in evidence))

    out: dict[str, Detection] = {}
    for key in ("project_name", "project_location", "owner"):
        d = Detection()
        distinct: list[tuple[str, str]] = []
        for v, where in cands.get(key, []):
            if not any(_same(v, u) for u, _ in distinct):
                distinct.append((v, where))
        lv = _clean((llm_ans or {}).get(key, "") or "")
        if lv and norm(lv) not in all_text and not any(_same(lv, u) for u, _ in distinct):
            lv = ""                                   # the model's answer is not in the file → ignore it
        if len(distinct) == 1:
            v, where = distinct[0]
            wheres = sorted({w for u, w in cands[key] if _same(u, v)})
            d.source = ", ".join(wheres[:2]) + (" …" if len(wheres) > 2 else "")
            if llm_ans is None:
                d.value, d.auto = v, True
                d.evidence = f"labelled in {', '.join(wheres[:3])}"
                d.note = "LLM cross-check not available — label match only"
            elif lv and _same(lv, v):
                d.value, d.auto = v, True
                d.evidence = f"labelled in {', '.join(wheres[:3])}; confirmed by the LLM"
            else:
                d.note = "label found but the LLM read it differently — not auto-filled"
                d.suggestions.append((v, f"labelled in {', '.join(wheres[:3])}"))
                if lv:
                    d.suggestions.append((lv, "read by the LLM"))
        elif len(distinct) > 1:
            d.note = "the BOQ states different values — not auto-filled"
            d.suggestions = [(v, f"labelled in {w}") for v, w in distinct]
        elif lv:
            d.note = "no explicit label — not auto-filled"
            d.suggestions.append((lv, "read by the LLM from the BOQ title area (unlabelled)"))
        out[key] = d
    return out


def load_field_config(path: Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))
