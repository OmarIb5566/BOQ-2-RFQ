"""Excel BOQ reader -> BOQDocument (rows, sections, items)."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Callable, Optional

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from ..models import BOQDocument, Row, SourceRef, TableScope
from ..text_utils import fmt_value, has_text, nfkc
from .structure import SECTION_RE, detect_header, build_structure, HeaderInfo
from ..text_utils import norm

_RANGE_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
_ROWS_RE = re.compile(r"\$?(\d+):\$?(\d+)")
# summary / collection sheets repeat division totals — they are not work packages
_SUMMARY_SHEET_RE = re.compile(r"summ|collection|grand\s*total|ملخص|اجمال|الاجمال|خلاصه|خلاصة", re.I)


def _print_area_cols(ws) -> Optional[tuple[int, int]]:
    pa = ws.print_area
    if not pa:
        return None
    pa = pa if isinstance(pa, str) else ",".join(pa)
    m = _RANGE_RE.search(pa)
    if not m:
        return None
    return column_index_from_string(m.group(1)), column_index_from_string(m.group(3))


def _print_title_rows(ws) -> Optional[tuple[int, int]]:
    pt = ws.print_title_rows
    if not pt:
        return None
    m = _ROWS_RE.search(str(pt))
    return (int(m.group(1)), int(m.group(2))) if m else None


def _hidden_cols(ws) -> set[int]:
    out = set()
    for d in ws.column_dimensions.values():
        if d.hidden and d.min and d.max:
            out.update(range(d.min, d.max + 1))
    return out


def read_excel(path: Path, doc_id: str, source_name: str, llm_header: Optional[Callable] = None,
               log: Callable[[str], None] = print) -> BOQDocument:
    wb = openpyxl.load_workbook(path, data_only=True)
    doc = BOQDocument(doc_id=doc_id, filename=source_name, source_path=str(path), work_path=str(path), kind="excel")
    sec_index = 0
    for s_idx, ws in enumerate(wb.worksheets):
        if ws.sheet_state != "visible":
            continue
        if _SUMMARY_SHEET_RE.search(norm(ws.title)):
            log(f"Sheet '{ws.title}': summary sheet — skipped.")
            continue
        pa_cols = _print_area_cols(ws)
        max_c = min(ws.max_column, pa_cols[1] if pa_cols else 80)
        first_c = pa_cols[0] if pa_cols else 1
        values: list[tuple] = []
        last_row, last_col = 0, first_c
        for r_i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=first_c, max_col=max_c,
                                               values_only=True), start=1):
            values.append(row)
            for c_off, v in enumerate(row):
                if v is not None and str(v).strip() != "":
                    last_row = r_i
                    last_col = max(last_col, first_c + c_off)
        if last_row == 0:
            continue
        values = values[:last_row]
        ncols = last_col - first_c + 1
        hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
        hidden_cols = _hidden_cols(ws)
        span: dict[int, int] = {}
        for mr in ws.merged_cells.ranges:
            if mr.max_row > mr.min_row:
                span[mr.min_row] = max(span.get(mr.min_row, mr.min_row), mr.max_row)

        def disp(r: int, c: int) -> str:
            v = values[r - 1][c - first_c] if c - first_c < len(values[r - 1]) else None
            return fmt_value(v)

        grid = []
        for r in range(1, min(last_row, 80) + 1):
            if r in hidden_rows:
                continue
            cells = [(c, disp(r, c)) for c in range(first_c, last_col + 1) if c not in hidden_cols]
            cells = [(c, t) for c, t in cells if t]
            grid.append((r, cells))
        header = detect_header(grid)
        if header is None and llm_header is not None:
            header = llm_header(ws.title, grid[:30])
        if header is None:
            if any(has_text(t) for _, cells in grid for _, t in cells):
                doc.warnings.append(f"Sheet '{ws.title}': no BOQ table header found — sheet skipped.")
            continue

        title_rows = _print_title_rows(ws)
        if title_rows and title_rows[0] <= header.best_row <= title_rows[1]:
            h_first, h_last = title_rows
        else:
            h_first, h_last = header.first_row, header.last_row

        fcols = header.field_columns
        mapped = {fcols[f] for f in ("item_no", "unit", "quantity", "rate", "total") if f in fcols}
        desc_col = fcols.get("description")

        prefix = f"X{s_idx}"
        rows: list[Row] = []
        header_uids: list[str] = []
        initial_title = ""
        for r in range(1, last_row + 1):
            if r in hidden_rows:
                continue
            cells = [disp(r, c) for c in range(first_c, last_col + 1)]
            ref = SourceRef(sheet=ws.title, row=r, span_to=span.get(r))
            if r <= h_last:
                if r >= h_first:
                    row = Row(uid=f"{prefix}R{r}", ref=ref, cells=cells, text=" ".join(t for t in cells if t), role="header")
                    doc.rows.append(row)
                    header_uids.append(row.uid)
                    for t in cells:
                        if t and SECTION_RE.search(norm(t)) and not initial_title:
                            initial_title = t
                continue
            fields = {f: disp(r, c) for f, c in fcols.items() if f != "description"}
            texts = []
            for c in range(first_c, last_col + 1):
                if c in mapped or c in hidden_cols:
                    continue
                t = disp(r, c)
                if t and (has_text(t) or c == desc_col):
                    texts.append(t)
            rows.append(Row(uid=f"{prefix}R{r}", ref=ref, cells=cells, fields=fields, text=" ".join(texts)))

        sections, items = build_structure(rows, prefix, initial_title or ws.title.strip(), ws.title, start_index=0)
        sec_index += len(sections)
        doc.rows.extend(rows)
        doc.sections.extend(sections)
        doc.items.extend(items)
        doc.scopes.append(TableScope(
            name=ws.title, kind="excel", header_row_uids=header_uids,
            header_labels={f: nfkc(v) for f, v in header.labels.items()},
            field_columns=fcols, first_col=first_c, last_col=last_col, last_row=last_row,
            hidden_rows=sorted(hidden_rows), hidden_cols=sorted(hidden_cols),
            rtl=bool(ws.sheet_view.rightToLeft),
        ))
        log(f"Sheet '{ws.title}': header rows {h_first}-{h_last}, columns "
            + ", ".join(f"{k}={get_column_letter(v)}" for k, v in fcols.items())
            + f"; {len(sections)} section(s), {len(items)} item(s).")
    if not doc.items:
        doc.warnings.append("No priced BOQ items were detected in the workbook.")
    return doc


def header_from_llm(llm, system: str):
    """Returns a callable used by read_excel when the keyword header detection fails."""
    from ..llm.prompts import HEADER_SCHEMA, HEADER_TASK

    def _call(sheet: str, grid) -> Optional[HeaderInfo]:
        lines = []
        for r, cells in grid:
            lines.append(f"{r}: " + " | ".join(f"{get_column_letter(c)}={t[:40]}" for c, t in cells))
        try:
            ans = llm.chat_json(system, HEADER_TASK.format(rows="\n".join(lines)), HEADER_SCHEMA)
        except Exception:  # noqa: BLE001
            return None
        fields = {}
        for f in ("item_no", "description", "unit", "quantity", "rate", "total"):
            col = (ans.get(f) or "").strip().upper()
            if re.fullmatch(r"[A-Z]{1,3}", col):
                fields[f] = column_index_from_string(col)
        # a real BOQ table has the description and its unit/quantity in different columns; the model
        # sometimes maps every field to the same column on summary / notes sheets
        desc = fields.get("description")
        if desc is None or not any(fields.get(f) not in (None, desc, fields.get("item_no"))
                                   for f in ("unit", "quantity")):
            return None
        first, last = int(ans["header_first_row"]), int(ans["header_last_row"])
        return HeaderInfo(best_row=last, first_row=first, last_row=max(first, last), field_columns=fields, labels={})

    return _call
