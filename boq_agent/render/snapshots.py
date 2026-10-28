"""Work-package snapshots ("cut-outs") as PDF.

A work package can hold items from several BOQ files of the project. Each file's rows are rendered as a
"part" with that file's own method, and the parts are joined into one PDF per work package:
* Excel source  -> Microsoft Excel hides every row that is not part of the WP and prints the sheet
                   (keeps the original formatting, logos, print titles and ALL columns).
                   Fallback without Excel: an HTML re-rendering of the same rows/columns.
* PDF source    -> the original page regions (header band + item rows, full table width) are
                   cropped and stacked onto new pages at vector quality.
"""
from __future__ import annotations

import contextlib
import html
import shutil
from collections import OrderedDict
from pathlib import Path
from typing import Callable

import pymupdf as fitz
from openpyxl.utils import get_column_letter

from ..config import Settings
from ..models import BOQDocument, Plan, SourceFile, WorkPackage
from ..planning.planner import snapshot_rows
from ..text_utils import fmt_value, safe_filename

Log = Callable[[str], None]


def _runs(nums: list[int]) -> list[tuple[int, int]]:
    out: list[tuple[int, int]] = []
    for n in sorted(set(nums)):
        if out and n == out[-1][1] + 1:
            out[-1] = (out[-1][0], n)
        else:
            out.append((n, n))
    return out


def wp_filename(wp: WorkPackage) -> str:
    return safe_filename(f"{wp.wp_id} {wp.name_en}") + ".pdf"


def _stamp(pdf_path: Path, wp: WorkPackage, settings: Settings) -> None:
    """Small WP label in the top margin of every page."""
    doc = fitz.open(pdf_path)
    arch = fitz.Archive(settings.font_dir)
    css = f"@font-face {{font-family: f; src: url({settings.font_file});}} * {{font-family: f; font-size: 7pt; color: #555;}}"
    label = html.escape(f"{wp.wp_id} | {wp.name_en}" + (f" | {wp.name_ar}" if wp.name_ar else ""))
    for i, page in enumerate(doc):
        r = page.rect
        page.insert_htmlbox(fitz.Rect(18, 3, r.width - 18, 16),
                            f"<div>{label} — page {i + 1}/{len(doc)}</div>", css=css, archive=arch)
    tmp = pdf_path.with_suffix(".tmp.pdf")
    doc.save(tmp, garbage=3, deflate=True)
    doc.close()
    shutil.move(tmp, pdf_path)


def _file_parts(doc: BOQDocument, wp: WorkPackage) -> list[tuple[SourceFile, list[str]]]:
    """The WP's snapshot rows grouped by source file, in project order."""
    rmap = doc.row_map()
    fmap = doc.file_map()
    parts: OrderedDict[str, list[str]] = OrderedDict()
    for uid in snapshot_rows(doc, wp):
        parts.setdefault(rmap[uid].ref.file, []).append(uid)
    return [(fmap[f], uids) for f, uids in parts.items() if f in fmap]


def _sheet_rows(doc: BOQDocument, uids: list[str]) -> "OrderedDict[str, list[int]]":
    rmap = doc.row_map()
    per_sheet: OrderedDict[str, set[int]] = OrderedDict()
    for uid in uids:
        ref = rmap[uid].ref
        per_sheet.setdefault(ref.sheet, set()).update(range(ref.row, (ref.span_to or ref.row) + 1))
    return OrderedDict((k, sorted(v)) for k, v in per_sheet.items())


def _header_rows(doc: BOQDocument, sheet: str, file: str) -> list[int]:
    rmap = doc.row_map()
    sc = doc.scope(sheet, file)
    return sorted(rmap[u].ref.row for u in sc.header_row_uids) if sc else []


# =============================================================================== Excel (MS Excel)
def _excel_com_part(xl, wb, doc: BOQDocument, f: SourceFile, uids: list[str], base: Path) -> list[Path]:
    parts = []
    for k, (sheet, rows) in enumerate(_sheet_rows(doc, uids).items()):
        sc = doc.scope(sheet, f.id)
        ws = wb.Worksheets(sheet)
        hdr = _header_rows(doc, sheet, f.id)
        last = max([sc.last_row] + rows)
        orig_hidden = set(sc.hidden_rows)
        show = (set(range(min(hdr), max(hdr) + 1)) if hdr else set()) | set(rows)
        ws.Rows(f"1:{last}").Hidden = True
        for a, b in _runs([r for r in show if r not in orig_hidden]):
            ws.Rows(f"{a}:{b}").Hidden = False
        with contextlib.suppress(Exception):
            xl.PrintCommunication = False
        ps = ws.PageSetup
        ps.PrintArea = f"${get_column_letter(sc.first_col)}$1:${get_column_letter(sc.last_col)}${last}"
        if hdr:
            ps.PrintTitleRows = f"${min(hdr)}:${max(hdr)}"
        ps.Zoom = False
        ps.FitToPagesWide = 1
        ps.FitToPagesTall = False
        with contextlib.suppress(Exception):
            xl.PrintCommunication = True
        ws.ResetAllPageBreaks()
        part = base.with_name(f"{base.stem}_{k}.pdf")
        ws.ExportAsFixedFormat(0, str(part.resolve()), 0, False, False)
        parts.append(part)
        ws.Rows(f"1:{last}").Hidden = False
        for a, b in _runs(sorted(orig_hidden)):
            ws.Rows(f"{a}:{b}").Hidden = True
    return parts


# =============================================================================== Excel (HTML fallback)
def _excel_html_part(doc: BOQDocument, f: SourceFile, uids: list[str], wp: WorkPackage, out: Path,
                     settings: Settings, wb_cache: dict) -> Path:
    import openpyxl

    if f.id not in wb_cache:
        wb_cache[f.id] = openpyxl.load_workbook(f.work_path, data_only=True)
    wb = wb_cache[f.id]
    arch = fitz.Archive(settings.font_dir)
    title = f"{wp.wp_id} — {wp.name_en} {wp.name_ar}" + (f" · {f.label}" if doc.multi_file else "")
    parts = [f"<h3>{html.escape(title)}</h3>"]
    max_cols = 1
    for sheet, rows in _sheet_rows(doc, uids).items():
        sc = doc.scope(sheet, f.id)
        ws = wb[sheet]
        hdr = _header_rows(doc, sheet, f.id)
        show = sorted(((set(range(min(hdr), max(hdr) + 1)) if hdr else set()) | set(rows)) - set(sc.hidden_rows))
        cols = [c for c in range(sc.first_col, sc.last_col + 1) if c not in set(sc.hidden_cols)]
        max_cols = max(max_cols, len(cols))
        parts.append(_sheet_table_html(ws, show, cols, set(hdr), sc.rtl or _looks_rtl(ws, show, cols)))
    css = (f"@font-face {{font-family: f; src: url({settings.font_file});}} "
           "* {font-family: f;} h3 {font-size: 11pt; margin: 0 0 6px 0;} "
           "table {border-collapse: collapse; width: 100%; margin-bottom: 10px;} "
           "td {border: 0.5px solid #777; padding: 2px 3px; font-size: 7pt; vertical-align: top;} "
           "tr.h td {background: #e8e8e8; font-weight: bold; text-align: center;}")
    story = fitz.Story(html="".join(parts), user_css=css, archive=arch)
    paper = fitz.paper_rect("a4-l" if max_cols > 5 else "a4")
    where = paper + (24, 30, -24, -24)
    writer = fitz.DocumentWriter(str(out))
    more = 1
    while more:
        dev = writer.begin_page(paper)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return out


def _looks_rtl(ws, rows, cols) -> bool:
    from ..text_utils import is_arabic
    sample = " ".join(str(ws.cell(r, c).value or "") for r in rows[:30] for c in cols)
    return is_arabic(sample)


def _sheet_table_html(ws, rows: list[int], cols: list[int], header_rows: set[int], rtl: bool) -> str:
    shown = set(rows)
    colset = set(cols)
    merge_top: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for mr in ws.merged_cells.ranges:
        first_vis_row = next((r for r in range(mr.min_row, mr.max_row + 1) if r in shown), None)
        first_vis_col = next((c for c in range(mr.min_col, mr.max_col + 1) if c in colset), None)
        if first_vis_row is None or first_vis_col is None:
            continue
        rs = sum(1 for r in range(mr.min_row, mr.max_row + 1) if r in shown)
        cs = sum(1 for c in range(mr.min_col, mr.max_col + 1) if c in colset)
        merge_top[(first_vis_row, first_vis_col)] = (rs, cs, mr.min_row, mr.min_col)
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if (r, c) != (first_vis_row, first_vis_col):
                    covered.add((r, c))
    widths = []
    for c in cols:
        w = ws.column_dimensions[get_column_letter(c)].width
        widths.append(w if w else 8.43)
    tot = sum(widths) or 1
    out = [f'<table dir="{"rtl" if rtl else "ltr"}"><colgroup>']
    out += [f'<col width="{100 * w / tot:.1f}%">' for w in widths]
    out.append("</colgroup>")
    for r in rows:
        out.append(f'<tr class="{"h" if r in header_rows else ""}">')
        for c in cols:
            if (r, c) in covered:
                continue
            attrs = ""
            src_r, src_c = r, c
            if (r, c) in merge_top:
                rs, cs, src_r, src_c = merge_top[(r, c)]
                attrs = f' rowspan="{rs}" colspan="{cs}"'
            cell = ws.cell(src_r, src_c)
            style = []
            if cell.font is not None and cell.font.b:
                style.append("font-weight:bold")
            fill = cell.fill
            if fill is not None and fill.fill_type == "solid":
                rgb = getattr(fill.fgColor, "rgb", None)
                if isinstance(rgb, str) and len(rgb) == 8 and rgb[2:] not in ("000000",):
                    style.append(f"background:#{rgb[2:]}")
            al = cell.alignment.horizontal if cell.alignment is not None else None
            if al in ("center", "centerContinuous"):
                style.append("text-align:center")
            elif al in ("left", "right"):
                style.append(f"text-align:{al}")
            text = html.escape(fmt_value(cell.value)).replace("\n", "<br>")
            out.append(f'<td{attrs} style="{";".join(style)}">{text}</td>')
        out.append("</tr>")
    out.append("</table>")
    return "".join(out)


# ================================================================================= PDF
def _pdf_part(doc: BOQDocument, f: SourceFile, uids: list[str], out: Path) -> Path:
    src = fitz.open(f.work_path)
    rmap = doc.row_map()
    file_rows = [r for r in doc.rows if r.ref.file == f.id]
    order = {r.uid: i for i, r in enumerate(file_rows)}
    scope = doc.scope("pdf", f.id)
    content_rows = [r for r in file_rows if r.role not in ("header", "blank")]
    pos = {r.uid: i for i, r in enumerate(content_rows)}
    uids = sorted((u for u in uids if u in pos), key=lambda u: order[u])
    # contiguous runs (in reading order, same page) -> one clip each
    clips: list[tuple[int, float, float]] = []
    prev_pos = None
    for u in uids:
        r = rmap[u]
        pg, (x0, y0, x1, y1) = r.ref.page, r.ref.bbox
        if clips and prev_pos is not None and pos[u] == prev_pos + 1 and clips[-1][0] == pg:
            p, a, b = clips[-1]
            clips[-1] = (p, min(a, y0), max(b, y1))
        else:
            clips.append((pg, y0, y1))
        prev_pos = pos[u]
    _compose(src, clips, scope, out)
    return out


def _compose(src, clips, scope, out: Path) -> None:
    dst = fitz.open()
    W, H = src[0].rect.width, src[0].rect.height
    margin, top = 20, 22
    page, y = None, 0.0

    def extent(pg):
        ext = scope.page_extents.get(pg) or scope.page_extents.get(str(pg))
        return ext or [0, src[pg].rect.width]

    def header_for(pg):
        hb = scope.page_headers.get(pg) or scope.page_headers.get(str(pg))
        if hb:
            return pg, hb
        keys = sorted(int(k) for k in scope.page_headers)
        prior = [k for k in keys if k <= pg] or keys
        if not prior:
            return None
        k = prior[-1]
        return k, scope.page_headers.get(k) or scope.page_headers.get(str(k))

    def new_page(pg):
        nonlocal page, y
        page = dst.new_page(width=W, height=H)
        y = top
        h = header_for(pg)
        if h:
            hp, hb = h
            x0, x1 = extent(hp)
            place(hp, fitz.Rect(x0, hb[1] - 3, x1, hb[3] + 3))

    def place(pg, clip):
        nonlocal y
        scale = min(1.0, (W - 2 * margin) / clip.width)
        h = clip.height * scale
        page.show_pdf_page(fitz.Rect(margin, y, margin + clip.width * scale, y + h), src, pg, clip=clip)
        y += h

    for pg, y0, y1 in clips:
        x0, x1 = extent(pg)
        clip = fitz.Rect(x0, y0 - 2, x1, y1 + 2)
        while clip.height > 0:
            if page is None:
                new_page(pg)
            scale = min(1.0, (W - 2 * margin) / clip.width)
            avail = (H - margin - y) / scale
            if avail < min(clip.height, 20):
                new_page(pg)
                continue
            part = fitz.Rect(clip.x0, clip.y0, clip.x1, clip.y0 + min(clip.height, avail))
            place(pg, part)
            clip = fitz.Rect(clip.x0, part.y1, clip.x1, clip.y1)
    if page is None:
        dst.new_page(width=W, height=H)
    dst.save(out, garbage=3, deflate=True)
    dst.close()


# ============================================================================ dispatch
def render_all(doc: BOQDocument, plan: Plan, out_dir: Path, settings: Settings, log: Log) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    tmp = out_dir / "_parts"
    tmp.mkdir(exist_ok=True)
    mode = settings.excel_renderer
    has_excel = any(f.kind == "excel" for f in doc.file_list())
    results: dict[str, Path] = {}
    with contextlib.ExitStack() as stack:
        xl = None
        if has_excel and mode in ("auto", "excel"):
            try:
                from ..excel_com import excel_app
                xl = stack.enter_context(excel_app())
            except Exception as e:  # noqa: BLE001
                if mode == "excel":
                    raise
                log(f"Excel rendering unavailable ({e}); using the built-in HTML renderer.")
        com_books: dict[str, object] = {}
        html_books: dict[str, object] = {}
        for wp in plan.work_packages:
            pieces: list[Path] = []
            for k, (f, uids) in enumerate(_file_parts(doc, wp)):
                base = tmp / f"{wp.wp_id}_{k}.pdf"
                if f.kind == "pdf":
                    pieces.append(_pdf_part(doc, f, uids, base))
                elif xl is not None:
                    if f.id not in com_books:
                        com_books[f.id] = xl.Workbooks.Open(str(Path(f.work_path).resolve()), UpdateLinks=0, ReadOnly=True)
                    pieces += _excel_com_part(xl, com_books[f.id], doc, f, uids, base)
                else:
                    pieces.append(_excel_html_part(doc, f, uids, wp, base, settings, html_books))
            out = out_dir / wp_filename(wp)
            merged = fitz.open()
            for p in pieces:
                with fitz.open(p) as d:
                    merged.insert_pdf(d)
            if merged.page_count == 0:
                merged.new_page()
            merged.save(out)
            merged.close()
            _stamp(out, wp, settings)
            results[wp.wp_id] = out
            log(f"{wp.wp_id}: {out.name} ({len(pieces)} part(s))")
        for wb in com_books.values():
            with contextlib.suppress(Exception):
                wb.Close(SaveChanges=False)
    shutil.rmtree(tmp, ignore_errors=True)
    return results
