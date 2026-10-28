"""Microsoft Excel automation (Windows only, optional)."""
from __future__ import annotations

import contextlib
import shutil
from pathlib import Path


def excel_available() -> bool:
    try:
        import pythoncom  # noqa: F401
        import win32com.client  # noqa: F401
    except ImportError:
        return False
    try:
        with excel_app() as xl:
            return bool(xl.Version)
    except Exception:  # noqa: BLE001
        return False


@contextlib.contextmanager
def excel_app():
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    xl = None
    try:
        xl = win32com.client.DispatchEx("Excel.Application")
        xl.Visible = False
        xl.DisplayAlerts = False
        xl.ScreenUpdating = False
        xl.EnableEvents = False
        xl.AskToUpdateLinks = False
        yield xl
    finally:
        if xl is not None:
            with contextlib.suppress(Exception):
                for wb in list(xl.Workbooks):
                    wb.Close(SaveChanges=False)
            with contextlib.suppress(Exception):
                xl.Quit()
        del xl
        pythoncom.CoUninitialize()


def convert_with_excel(src: Path, dst: Path) -> Path:
    """Save any Excel-readable workbook (.xls/.xlsb/...) as .xlsx."""
    with excel_app() as xl:
        wb = xl.Workbooks.Open(str(src.resolve()), UpdateLinks=0, ReadOnly=True)
        wb.SaveAs(str(dst.resolve()), FileFormat=51)
        wb.Close(SaveChanges=False)
    return dst


def convert_with_xlrd(src: Path, dst: Path) -> Path:
    """Values-only .xls -> .xlsx conversion when Excel is not installed (formatting is lost)."""
    import openpyxl
    import xlrd

    book = xlrd.open_workbook(str(src), formatting_info=False)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v not in ("", None):
                    ws.cell(row=r + 1, column=c + 1, value=v)
        for (r0, r1, c0, c1) in getattr(sh, "merged_cells", []):
            ws.merge_cells(start_row=r0 + 1, end_row=r1, start_column=c0 + 1, end_column=c1)
    wb.save(dst)
    return dst


def prepare_workbook(src: Path, work_dir: Path, log=print, stem: str = "source") -> Path:
    """Copy (and if needed convert) the uploaded workbook into the run folder; returns a .xlsx/.xlsm path.
    `stem` must be unique per project: Excel cannot open two workbooks with the same file name at once."""
    ext = src.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        dst = work_dir / f"{stem}{ext}"
        shutil.copy2(src, dst)
        return dst
    dst = work_dir / f"{stem}.xlsx"
    try:
        log(f"Converting {src.name} to .xlsx with Microsoft Excel …")
        return convert_with_excel(src, dst)
    except Exception as e:  # noqa: BLE001
        if ext == ".xls":
            log(f"Excel conversion unavailable ({e}); falling back to xlrd (values only).")
            return convert_with_xlrd(src, dst)
        raise RuntimeError(f"Cannot read {src.name}: Microsoft Excel is required for {ext} files ({e})") from e
